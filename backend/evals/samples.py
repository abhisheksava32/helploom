"""Builders for the binary sample documents (DOCX and PDF).

The Markdown, HTML and text samples in evals/kb/ are written by hand. The Word and PDF
samples are built from the text below so their content stays reviewable in a diff;
rebuild them after editing with:

    uv run python -m evals.samples

make_pdf and make_docx are also used by the test suites to produce real files of
both formats without fixtures checked in as opaque binaries.
"""

from __future__ import annotations

import datetime as dt
import io
import textwrap
from collections.abc import Sequence
from pathlib import Path

KB_DIR = Path(__file__).resolve().parent / "kb"
AUTHOR = "Kettlewick Coffee Co."


def _pdf_escape(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: Sequence[Sequence[str]], *, wrap: int = 88) -> bytes:
    """A small, valid PDF with a real text layer.

    Each page is a list of paragraphs; a paragraph starting with "# " is drawn as a
    bold heading. Uses the standard Helvetica fonts, so no font embedding is needed.
    """
    objects: list[bytes] = []

    def add(body: str | bytes) -> int:
        objects.append(body.encode("latin-1") if isinstance(body, str) else body)
        return len(objects)

    catalog = add("")  # filled in once the page tree exists
    pages_id = add("")
    regular = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")

    page_ids: list[int] = []
    for paragraphs in pages:
        lines = ["BT", "15 TL", "72 740 Td"]
        for paragraph in paragraphs:
            if paragraph.startswith("# "):
                lines += ["/F2 13 Tf", f"({_pdf_escape(paragraph[2:])}) Tj T*"]
            else:
                lines.append("/F1 10.5 Tf")
                lines += [f"({_pdf_escape(row)}) Tj T*" for row in textwrap.wrap(paragraph, wrap)]
            lines.append("T*")
        lines.append("ET")
        stream = "\n".join(lines).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        page_ids.append(
            add(
                f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 612 792] "
                f"/Resources << /Font << /F1 {regular} 0 R /F2 {bold} 0 R >> >> /Contents {content} 0 R >>"
            )
        )
    kids = " ".join(f"{page} 0 R" for page in page_ids)
    objects[pages_id - 1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("latin-1")
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode("latin-1")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, catalog, xref)
    )
    return out.getvalue()


def make_docx(blocks: Sequence[tuple[str, object]]) -> bytes:
    """A Word document from (kind, value) blocks: ("h1".."h3", text), ("p", text),
    ("bullet", text) or ("table", rows)."""
    import docx

    document = docx.Document()
    properties = document.core_properties
    # The default template ships with its own author and comment; replace them.
    properties.author = AUTHOR
    properties.last_modified_by = AUTHOR
    properties.comments = ""
    properties.title = ""
    fixed = dt.datetime(2026, 9, 1, 9, 0, 0)
    properties.created = fixed
    properties.modified = fixed
    properties.revision = 1

    for kind, value in blocks:
        if kind in ("h1", "h2", "h3"):
            document.add_heading(str(value), level=int(kind[1]))
        elif kind == "p":
            document.add_paragraph(str(value))
        elif kind == "bullet":
            document.add_paragraph(str(value), style="List Bullet")
        elif kind == "table":
            rows = [list(map(str, row)) for row in value]  # type: ignore[attr-defined]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for row_cells, row in zip(table.rows, rows, strict=True):
                for cell, text in zip(row_cells.cells, row, strict=True):
                    cell.text = text
        else:
            raise ValueError(f"Unknown block kind {kind!r}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


EQUIPMENT_CARE: list[tuple[str, object]] = [
    ("h1", "Equipment care and warranty"),
    (
        "p",
        "This guide covers everyday care for the Kettlewick Mill M2 burr grinder and the Pourwell "
        "gooseneck kettle, the meaning of the grinder's error codes, and how our warranty works.",
    ),
    ("h2", "Cleaning the Mill M2 grinder"),
    (
        "p",
        "Brush the burrs every 2 weeks with the cleaning brush stored in the base. Unplug the grinder, "
        "empty the hopper, twist the upper burr counter clockwise to lift it out and brush away the "
        "coffee fines from both burrs.",
    ),
    (
        "p",
        "Once a month, run a dose of grinder cleaning tablets through the burrs to absorb coffee oils. "
        "Do not use rice: it is harder than coffee, can chip the burrs and voids the warranty.",
    ),
    ("bullet", "Never wash the burrs with water; moisture makes them rust."),
    ("bullet", "Wipe the hopper and the grounds bin with a dry cloth."),
    ("bullet", "Replacement burr sets cost $24 and take five minutes to fit."),
    ("h2", "Descaling the Pourwell kettle"),
    (
        "p",
        "Descale the kettle once a month, or every two weeks in hard water areas. Fill it with 1 litre "
        "of water mixed with 1 tablespoon of citric acid, bring it to a boil, leave it for 20 minutes, "
        "then pour it out and rinse twice with fresh water.",
    ),
    (
        "p",
        "The Hold mode keeps water at the set temperature for 60 minutes, then the kettle switches off "
        "on its own to save energy.",
    ),
    ("h2", "Mill M2 error codes"),
    (
        "table",
        [
            ["Code", "Meaning", "What to do"],
            ["E1", "Motor jam", "Unplug, remove the hopper and clear any beans stuck between the burrs."],
            ["E2", "Motor overheating", "Let the grinder cool down for 20 minutes before grinding again."],
            ["E3", "Hopper not seated", "Lift the hopper and turn it until it clicks into place."],
        ],
    ),
    ("p", "If a code keeps coming back after these steps, contact support with the serial number."),
    ("h2", "Warranty coverage"),
    (
        "p",
        "Electrical equipment such as the Mill M2 grinder and the Pourwell kettle has a 2 year warranty "
        "from the date of purchase. Accessories such as drippers, carafes and filter holders have a 1 "
        "year warranty. Burrs are wear parts and are covered for 1 year.",
    ),
    (
        "p",
        "The warranty does not cover commercial use, damage from grinding anything other than coffee "
        "beans, or repairs by anyone other than Kettlewick or an authorised service partner.",
    ),
    ("h2", "Making a warranty claim"),
    (
        "p",
        "To make a claim, email support with the serial number and your proof of purchase. The serial "
        "number is printed under the base and starts with KW-. We reply within 2 business days with a "
        "repair, a replacement or a prepaid label to send the item in.",
    ),
]

REWARDS_PROGRAM: list[list[str]] = [
    [
        "# Kettlewick Rewards",
        "Kettlewick Rewards is our free loyalty program. Every account joins automatically, and points "
        "are added to your account as soon as an order ships.",
        "# Earning Points",
        "You earn 1 point for every $1 you spend on coffee, equipment and accessories. Points are not "
        "earned on gift cards, shipping costs or taxes.",
        "You also earn a birthday bonus of 50 points every year when you add your birthday in Settings "
        "> Profile.",
        "# Referring Friends",
        "Share your referral link from Account > Rewards. Your friend gets 15% off their first order, "
        "and you earn 200 points when that order ships.",
    ],
    [
        "# Redeeming Points",
        "Every 100 points can be turned into a $5 reward at checkout. Rewards can be combined with "
        "free shipping but not with promo codes, and they cannot be exchanged for cash.",
        "# Membership Tiers",
        "Your tier depends on the points you earned in the last 12 months. Steep members have 0 to 299 "
        "points, Bloom members have 300 to 799 points and Roast members have 800 points or more.",
        "Bloom members get early access to limited coffees. Roast members also get free express "
        "shipping on every order and an invitation to our annual cupping event.",
        "# Points Expiry",
        "Points expire after 12 months without any account activity. Any purchase or redeemed reward "
        "counts as activity and keeps all of your points.",
    ],
]


def write_samples(directory: Path = KB_DIR) -> list[Path]:
    targets = {
        directory / "equipment-care-and-warranty.docx": make_docx(EQUIPMENT_CARE),
        directory / "rewards-program.pdf": make_pdf(REWARDS_PROGRAM),
    }
    for path, data in targets.items():
        path.write_bytes(data)
    return list(targets)


if __name__ == "__main__":
    for path in write_samples():
        print(f"Wrote {path.relative_to(KB_DIR.parent.parent)}")
